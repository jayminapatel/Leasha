"""Layer 2 acceptance tests.

The criteria from BUILD_SPEC_V2.md, verbatim:

  1. A fixtures folder with one healthy and one deliberately corrupt file of
     each type.
  2. Every healthy fixture yields non-empty text and plausible chunk counts.
  3. Every corrupt fixture yields ERR_FILE_CORRUPT with SKIP_CONTINUE - and the
     run continues.
  4. A password-protected PDF and a locked-open XLSX both skip cleanly.
  5. PST extraction over a small test archive preserves conversation grouping.
  6. The live Outlook mailbox is enumerated alongside .pst stores, and closing
     Outlook mid-run yields ERR_OUTLOOK_BUSY rather than failing the run.
  7. A OneDrive folder with Files On-Demand indexes pinned files and skips
     placeholders, and no placeholder is hydrated.
  8. Chunk overlap verified: no text lost at boundaries.

**Criteria 5 and 6** are covered through a fake MAPI session. The COM calls live
behind `Win32ComSession`; the walk, the conversation grouping and the error
handling above it are ordinary code and are tested here on any machine. What
remains untested is that the adapter drives real Outlook - one clearly marked
`xfail(run=False)`, rather than the whole criterion.

Criterion 7 is covered at the attribute level: `winfs.is_cloud_placeholder`
accepts injected attribute bits, so placeholder *detection* is proved off
Windows. That no placeholder is hydrated must still be confirmed by watching the
sync client on a real machine, exactly as the spec says.
"""

from __future__ import annotations

from itertools import pairwise
from pathlib import Path

import pytest

from app.core.errors import ActionType, AppErrorException
from app.core.winfs import (
    FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS,
    describe_placeholder,
    is_cloud_placeholder,
)
from app.extract import chunk_document, extract
from app.extract.chunker import chunk_text
from app.extract.email_pst import walk_session

HEALTHY_OF_EACH_TYPE = {
    "pdf": "pdf/healthy.pdf",
    "docx": "office/healthy.docx",
    "xlsx": "office/healthy.xlsx",
    "pptx": "office/healthy.pptx",
    "eml": "email/thread_root.eml",
    "txt": "plaintext/utf8.txt",
    "csv": "plaintext/table.csv",
}

BAD_OF_EACH_TYPE = {
    "pdf-corrupt": "pdf/truncated.pdf",
    "pdf-encrypted": "pdf/encrypted.pdf",
    "pdf-scanned": "pdf/scanned.pdf",
    "pdf-lying-extension": "corrupt/lies.pdf",
    "docx-empty": "corrupt/empty.docx",
    "xlsx-not-a-package": "corrupt/notoffice.xlsx",
    "unsupported": "corrupt/unsupported.xyz",
    "txt-binary": "plaintext/binary.log",
    "eml-no-content": "email/empty_body.eml",
}


# --- 1 ----------------------------------------------------------------------

def test_fixture_corpus_covers_every_type(fixture_root: Path) -> None:
    for name in (*HEALTHY_OF_EACH_TYPE.values(), *BAD_OF_EACH_TYPE.values()):
        assert (fixture_root / name).exists(), f"missing fixture: {name}"


# --- 2 ----------------------------------------------------------------------

@pytest.mark.parametrize("name", list(HEALTHY_OF_EACH_TYPE.values()), ids=list(HEALTHY_OF_EACH_TYPE))
def test_healthy_fixtures_yield_text_and_plausible_chunks(fixture_root: Path, name: str) -> None:
    documents = list(extract(fixture_root / name))
    assert len(documents) == 1

    document = documents[0]
    assert len(document.text.strip()) > 20

    chunks = chunk_document(document)
    assert chunks, "text but no chunks"
    assert [chunk.ordinal for chunk in chunks] == list(range(len(chunks)))
    for chunk in chunks:
        assert chunk.text.strip()
        assert document.text[chunk.char_start:chunk.char_end] == chunk.text


# --- 3 ----------------------------------------------------------------------

@pytest.mark.parametrize("name", list(BAD_OF_EACH_TYPE.values()), ids=list(BAD_OF_EACH_TYPE))
def test_bad_fixtures_skip_and_continue(fixture_root: Path, name: str) -> None:
    with pytest.raises(AppErrorException) as caught:
        list(extract(fixture_root / name))

    error = caught.value.error
    assert error.action_type is ActionType.SKIP_CONTINUE
    assert error.suggestion.strip()
    assert not error.is_fatal


def test_a_batch_survives_every_bad_file_in_the_corpus(fixture_root: Path) -> None:
    """The layer's whole purpose: one bad file never halts a 100GB run."""
    everything = sorted(
        path
        for path in fixture_root.rglob("*")
        if path.is_file() and path.suffix not in {".py", ".md"} and path.name != ".gitkeep"
    )
    assert len(everything) >= 15

    indexed, skipped = 0, 0
    for path in everything:
        try:
            for document in extract(path):
                chunk_document(document)
                indexed += 1
        except AppErrorException:
            skipped += 1

    assert indexed >= 7
    assert skipped >= 8
    assert indexed + skipped == len(everything)


# --- 4 ----------------------------------------------------------------------

def test_password_protected_pdf_skips_cleanly(fixture_root: Path) -> None:
    with pytest.raises(AppErrorException) as caught:
        list(extract(fixture_root / "pdf/encrypted.pdf"))
    error = caught.value.error
    assert error.code == "ERR_FILE_CORRUPT"
    assert error.action_type is ActionType.SKIP_CONTINUE
    assert "password" in error.details.lower()


def test_locked_file_maps_to_the_locked_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """A file held open by Excel raises PermissionError on Windows.

    POSIX will not reproduce that lock, so the *mapping* is what is proved here:
    a PermissionError must become ERR_FILE_LOCKED ('close the program holding
    it'), never ERR_FILE_CORRUPT ('this file is damaged'). Sending someone to
    repair a perfectly good file is the failure this guards against.
    """
    target = tmp_path / "budget.xlsx"
    target.write_bytes(b"irrelevant - load_workbook is patched")

    import openpyxl

    def refuse(*_args, **_kwargs):
        raise PermissionError(32, "The process cannot access the file")

    monkeypatch.setattr(openpyxl, "load_workbook", refuse)

    with pytest.raises(AppErrorException) as caught:
        list(extract(target))

    error = caught.value.error
    assert error.code == "ERR_FILE_LOCKED"
    assert error.action_type is ActionType.SKIP_CONTINUE
    assert "close" in error.suggestion.lower()


# --- 5 and 6 ----------------------------------------------------------------
#
# The COM calls live behind `Win32ComSession`; the walk, the grouping and the
# error handling are ordinary code driven here through a fake implementing the
# same duck types. That covers the logic on any machine. What it cannot cover is
# that the adapter drives real Outlook correctly - see the residual xfail below,
# which is the honest remainder rather than the whole criterion.

def test_pst_preserves_conversation_grouping() -> None:
    """Criterion 5. A decision is rarely in one message, so replies must group."""
    from tests.unit.test_email_pst import FakeFolder, FakeSession, FakeStore, message

    inbox = FakeFolder("Inbox", items=[
        message("m1", "Contract terms", "Proposing 30 days.", conversation="thread-x"),
        message("m2", "Re: Contract terms", "Agreed.", conversation="thread-x"),
        message("m3", "Unrelated", "Lunch?", conversation="thread-y"),
    ])
    root = FakeFolder("Top", children=[inbox])
    session = FakeSession(FakeStore("archive", root, file_path=r"D:\a.pst"))

    threads: dict[str, list[str]] = {}
    for document in walk_session(session):
        threads.setdefault(document.meta["conversation"], []).append(document.meta["subject"])

    assert sorted(threads["thread-x"]) == ["Contract terms", "Re: Contract terms"]
    assert threads["thread-y"] == ["Unrelated"]


def test_live_mailbox_is_enumerated_alongside_archives() -> None:
    """Criterion 6, first half: the live mailbox is walked *as well as* the
    .pst stores, not instead of them."""
    from tests.unit.test_email_pst import FakeFolder, FakeSession, FakeStore, message

    live = FakeStore("Mailbox - Jaymin", FakeFolder("Top", children=[
        FakeFolder("Inbox", items=[message("live-1", "From the live mailbox")])
    ]), file_path=r"C:\cache.ost", is_live=True, cached_only=True)
    archive = FakeStore("2007", FakeFolder("Top", children=[
        FakeFolder("Inbox", items=[message("arch-1", "From the archive")])
    ]), file_path=r"D:\2007.pst")

    documents = list(walk_session(FakeSession(live, archive)))
    assert {d.meta["subject"] for d in documents} == {
        "From the live mailbox", "From the archive",
    }
    cached = {d.meta["store_name"]: d.meta["store_cached_only"] for d in documents}
    assert cached["Mailbox - Jaymin"] is True, "the cache gap must be recorded, not hidden"
    assert cached["2007"] is False


def test_closing_outlook_mid_run_yields_err_outlook_busy() -> None:
    """Criterion 6, second half. Losing a folder must not lose the run - by the
    time Outlook closes, thousands of messages may already have been read."""
    from app.extract.email_pst import drain_busy_folders
    from tests.unit.test_email_pst import FakeFolder, FakeSession, FakeStore, message

    drain_busy_folders()
    done = FakeFolder("Inbox", items=[message("a"), message("b")])
    closed = FakeFolder("Archive", items=[message("c")])
    closed.raises_on_items = True
    remaining = FakeFolder("Projects", items=[message("d")])
    root = FakeFolder("Top", children=[done, closed, remaining])

    documents = list(walk_session(FakeSession(FakeStore("s", root, file_path=r"D:\a.pst"))))
    assert len(documents) == 3, "only the unreachable folder was lost"

    busy = drain_busy_folders()
    assert [b.code for b in busy] == ["ERR_OUTLOOK_BUSY"]
    assert busy[0].action_type is ActionType.SKIP_CONTINUE
    assert "Outlook open" in busy[0].suggestion


@pytest.mark.xfail(
    run=False,
    reason=(
        "The residue: that Win32ComSession drives real Outlook correctly. Needs Windows "
        "with Outlook running and a real .pst. Everything above it is covered."
    ),
)
def test_win32com_adapter_against_real_outlook() -> None:
    raise AssertionError("run manually: app.cli extract --mailbox")


# --- 7 ----------------------------------------------------------------------

def test_cloud_placeholders_are_detected_before_anything_reads_them(tmp_path: Path) -> None:
    """Detection is by attribute bits, so it never opens the file - which is the
    point: opening a placeholder is what triggers the download."""
    pinned = tmp_path / "pinned.txt"
    pinned.write_text("locally available", encoding="utf-8")

    assert not is_cloud_placeholder(pinned, attributes=0)
    assert is_cloud_placeholder(pinned, attributes=FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS)
    assert describe_placeholder(FILE_ATTRIBUTE_RECALL_ON_DATA_ACCESS) == "RECALL_ON_DATA_ACCESS"


def test_cloud_only_error_explains_the_tradeoff() -> None:
    from app.core.errors import make_error

    error = make_error("ERR_CLOUD_ONLY", "index.walker", path="D:\\OneDrive\\report.docx")
    assert error.action_type is ActionType.SKIP_CONTINUE
    assert "Always keep on this device" in error.suggestion


# --- 8 ----------------------------------------------------------------------

def test_no_text_is_lost_at_chunk_boundaries(fixture_root: Path) -> None:
    """Across every healthy fixture, at a chunk size small enough to force many
    boundaries, every word must survive into at least one chunk."""
    import re

    word = re.compile(r"\S+")

    for name in HEALTHY_OF_EACH_TYPE.values():
        document = next(iter(extract(fixture_root / name)))
        chunks = chunk_text(document.text, target_tokens=24, overlap_tokens=6)
        spans = [(chunk.char_start, chunk.char_end) for chunk in chunks]

        missing = [
            match.group()
            for match in word.finditer(document.text)
            if not any(start <= match.start() and match.end() <= end for start, end in spans)
        ]
        assert not missing, f"{name}: {len(missing)} word(s) lost at a boundary: {missing[:5]}"


def test_overlap_actually_overlaps(fixture_root: Path) -> None:
    document = next(iter(extract(fixture_root / "pdf/healthy.pdf")))
    chunks = chunk_text(document.text, target_tokens=20, overlap_tokens=8, min_chunk_chars=0)
    assert len(chunks) > 1
    for previous, following in pairwise(chunks):
        assert following.char_start < previous.char_end
