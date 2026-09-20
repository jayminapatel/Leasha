r"""The `.pptx` fast path (`app/extract/ooxml_pptx.py`).

Checked against python-pptx itself on a deck python-pptx builds, so the shape the
fast reader must reproduce - slide order, notes, tables, groups, chart labels - is
whatever the reference library reports, not what this file's author remembers.
"""

from __future__ import annotations

from pathlib import Path

import pytest

import app.extract  # noqa: F401
from app.core.errors import AppErrorException
from app.extract.base import REGISTRY, extract
from app.extract.ooxml_pptx import PptxUnreadable, pptx_slides

pptx = pytest.importorskip("pptx")


def build_deck(path: Path) -> Path:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = pptx.Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Boiler review"
    first.placeholders[1].text = "Licence 12400 is current"
    first.notes_slide.notes_text_frame.text = "Remember to mention the Leeds site"

    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Costs"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Burner"
    table.cell(1, 1).text = "4200"
    group = second.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1))
    box.text_frame.text = "Inside a group"

    third = deck.slides.add_slide(deck.slide_layouts[5])
    third.shapes.title.text = "Output"
    data = CategoryChartData()
    data.categories = ["North", "South"]
    data.add_series("Throughput", (1, 2))
    third.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(3), data)
    deck.save(str(path))
    return path


def test_slides_come_back_in_show_order_with_notes_kept_apart(tmp_path: Path) -> None:
    slides = pptx_slides(build_deck(tmp_path / "deck.pptx"))
    assert len(slides) == 3
    assert "Boiler review" in slides[0][0] and "Licence 12400 is current" in slides[0][0]
    assert slides[0][1] == "Remember to mention the Leeds site"
    assert slides[1][1] == ""


def test_tables_groups_and_chart_labels_are_read(tmp_path: Path) -> None:
    slides = pptx_slides(build_deck(tmp_path / "deck.pptx"))
    assert "Item\tAmount" in slides[1][0] and "Burner\t4200" in slides[1][0]
    assert "Inside a group" in slides[1][0]
    assert "North South" in slides[2][0] and "Throughput" in slides[2][0]


def test_the_document_has_the_same_labels_as_the_python_pptx_path(tmp_path: Path) -> None:
    documents = list(extract(build_deck(tmp_path / "deck.pptx")))
    assert documents[0].meta["read_by"] == "zip+xml"
    labels = [segment.label for segment in documents[0].segments]
    assert labels[:2] == ["Slide 1", "Slide 1 speaker notes"]
    assert [segment.page for segment in documents[0].segments][:3] == [1, 1, 2]


def test_a_slide_show_and_a_template_are_readable(tmp_path: Path) -> None:
    for extension in (".ppsx", ".potx"):
        assert extension in REGISTRY
        path = build_deck(tmp_path / "deck.pptx").rename(tmp_path / f"copy{extension}")
        assert "Licence 12400" in "\n".join(d.text for d in extract(path))


def test_something_that_is_not_a_package_is_refused_so_the_old_path_reports_it(tmp_path: Path) -> None:
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"not a zip")
    with pytest.raises(PptxUnreadable):
        pptx_slides(path)
    with pytest.raises(AppErrorException) as caught:
        list(extract(path))
    assert caught.value.error.code == "ERR_FILE_CORRUPT"
