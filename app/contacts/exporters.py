"""Exporting contacts: vCard, CSV (Google, Outlook, plain), Excel and PDF.

Layer: L7

Order 1k 6j. Pure functions - cards in, text or a file out; the Contacts page
chooses the cards (everything, the selection or the filter) and the file.

**Round trips.** Exported, then imported into an empty book (`importers`), a card
has the same `card.units(c, LAYOUT_CAPS[layout])` for every field the layout can
hold. What each one holds:

* `vcard` (3.0 and 4.0) and `plain` CSV: everything - every scalar, every kind,
  groups, favourite and the photo. vCard keeps labels in lower case; the
  favourite star travels as `X-LEASHA-FAVOURITE:1`, which other programs ignore.
  Instant-message addresses are written as `IMPP;TYPE=<service>:<service>:<name>`
  so their service survives (`vcard.dumps` alone would lose it in 4.0).
* `xlsx`: the plain layout without the photo (a cell holds 32,767 characters).
* `google`: all scalars; email, phone, address, website, date, related person and
  custom field; groups and the star. No photo, social profile or IM address, no
  "main" flag. A birthday goes in Google's Birthday column and every other date
  is an Event. Several values of one kind go in numbered columns.
* `outlook`: all scalars (Outlook's "Title" column is the name prefix); up to 3
  email addresses, no labels on them; phones in Outlook's named slots (two home,
  two business, mobile, car, other, pager, primary, three faxes - a number with
  no free slot is not written); one address each for business, home and other;
  two web pages and one IM address, without labels; a birthday and an anniversary;
  spouse, manager and assistant as related people; groups as Categories. No
  photo, star, main flag, social profile or custom field.

`LAYOUT_CAPS[layout]` is a `card.Capabilities` for each, which is what the export
dialog uses to say what will not be saved, and what the tests compare with.
"""

from __future__ import annotations

import base64
import csv
import io
import re
from pathlib import Path
from typing import Iterable

from app.contacts import _layouts as L
from app.contacts import _pdf as P
from app.contacts import vcard
from app.contacts.card import Card, Field, display_name
from app.contacts._layouts import LAYOUT_CAPS, MULTI, label_out
from app.contacts._pdf import NothingToPrint, PdfLayout, Placement

__all__ = ["LAYOUT_CAPS", "export_vcard", "vcard_text", "export_csv", "write_csv", "export_xlsx",
           "export_pdf_list", "export_pdf_cards", "csv_table", "FAVOURITE_MARK"]

FAVOURITE_MARK = "X-LEASHA-FAVOURITE"


def _live(card: Card) -> list[Field]:
    return [f for f in card.fields if f.value.strip() or any(v for v in f.data.values())]


# --------------------------------------------------------------------------
# vCard
# --------------------------------------------------------------------------

def vcard_text(card: Card, version: str = "3.0") -> str:
    """One card as a vCard, with the favourite star and instant-message services kept."""
    if version not in ("3.0", "4.0"):
        raise ValueError("vCard version must be 3.0 or 4.0.")
    base = card.copy()
    base.fields = [f for f in _live(card) if f.kind != "im"]
    text = vcard.dumps(base, version=version)
    extra: list[str] = []
    for f in _live(card):
        if f.kind == "im":
            label = re.sub(r"[,;:\"\\]", "", f.label).strip()
            scheme = re.sub(r"[^a-z0-9+.-]", "", f.label.lower()) or "im"
            tparam = f";TYPE={label}" if label else ""
            extra.append(vcard._fold(f"IMPP{tparam}:{scheme}:{vcard._esc(f.value)}"))
    if card.favourite:
        extra.append(f"{FAVOURITE_MARK}:1")
    if extra:
        head, tail = text.rsplit("END:VCARD", 1)
        text = head + "\r\n".join(extra) + "\r\nEND:VCARD" + tail
    return text


def _file_name(card: Card, taken: set[str]) -> str:
    base = re.sub(r"[^\w .'-]", "", display_name(card), flags=re.UNICODE).strip(" .") or "contact"
    name, n = base, 1
    while f"{name}.vcf".lower() in taken:
        n += 1
        name = f"{base} ({n})"
    taken.add(f"{name}.vcf".lower())
    return f"{name}.vcf"


def export_vcard(cards: Iterable[Card], version: str = "3.0", one_file: bool = True):
    """`one_file`: one string holding every card (`""` for none). Otherwise a list of
    `(file name, text)`, one per contact, with names that are safe on Windows and unique."""
    cards = list(cards)
    if one_file:
        return "".join(vcard_text(c, version) for c in cards)
    if version not in ("3.0", "4.0"):
        raise ValueError("vCard version must be 3.0 or 4.0.")
    taken: set[str] = set()
    return [(_file_name(c, taken), vcard_text(c, version)) for c in cards]


# --------------------------------------------------------------------------
# CSV tables: (headers, rows)
# --------------------------------------------------------------------------

def _by_kind(card: Card, kind: str) -> list[Field]:
    return [f for f in _live(card) if f.kind == kind]


def _address_text(f: Field) -> str:
    return f.value or ", ".join(f.data.get(k, "") for k in ("street", "city", "region", "postcode", "country")
                                if f.data.get(k))


def _plain_table(cards: list[Card], with_photo: bool = True) -> tuple[list[str], list[list[str]]]:
    counts = {k: max([len(_by_kind(c, k)) for c in cards] + [1 if k in ("email", "phone", "address", "url", "date") else 0])
              for k, _ in L.PLAIN_KINDS}
    headers = [h for h, _ in L.PLAIN_SCALARS] + ["Groups", "Favourite"] + (["Photo"] if with_photo else [])
    for kind, title in L.PLAIN_KINDS:
        for n in range(1, counts[kind] + 1):
            headers += [f"{title} {n} Label", f"{title} {n} Value"]
            if kind == "address":
                headers += [f"{title} {n} {p}" for p in ("Street", "City", "Region", "Postcode", "Country")]
            headers.append(f"{title} {n} Primary")
    rows = []
    for c in cards:
        row = [c.get(name) for _, name in L.PLAIN_SCALARS]
        row += [MULTI.join(g for g in c.groups if g.strip()), "yes" if c.favourite else ""]
        if with_photo:
            row.append(f"data:{c.photo.mime};base64,{base64.b64encode(c.photo.data).decode('ascii')}"
                       if c.photo is not None else "")
        for kind, _title in L.PLAIN_KINDS:
            have = _by_kind(c, kind)
            for n in range(counts[kind]):
                f = have[n] if n < len(have) else None
                row += [f.label if f else "", (f.value if f else "")]
                if kind == "address":
                    d = f.data if f else {}
                    row += [d.get(p, "") for p in ("street", "city", "region", "postcode", "country")]
                row.append("yes" if f and f.primary else "")
        rows.append(row)
    return headers, rows


def _google_table(cards: list[Card]) -> tuple[list[str], list[list[str]]]:
    def events(c: Card) -> list[Field]:
        dates = _by_kind(c, "date")
        first = next((f for f in dates if f.label.strip().lower() == "birthday"), None)
        return [f for f in dates if f is not first]

    spec = [("E-mail", "email", ["Label", "Value"]),
            ("Phone", "phone", ["Label", "Value"]),
            ("Address", "address", ["Label", "Formatted", "Street", "City", "PO Box", "Region", "Postal Code",
                                    "Country", "Extended Address"]),
            ("Website", "url", ["Label", "Value"]),
            ("Event", "event", ["Label", "Value"]),
            ("Relation", "related", ["Label", "Value"]),
            ("Custom Field", "custom", ["Label", "Value"])]

    def of(c: Card, kind: str) -> list[Field]:
        return events(c) if kind == "event" else _by_kind(c, kind)

    counts = {k: max([len(of(c, k)) for c in cards] + [1 if k in ("email", "phone") else 0]) for _, k, _ in spec}
    headers = ["First Name", "Middle Name", "Last Name", "Name Prefix", "Name Suffix", "Nickname",
               "Organization Name", "Organization Title", "Organization Department", "Birthday", "Notes",
               "Photo", "Labels"]
    for title, kind, subs in spec:
        for n in range(1, counts[kind] + 1):
            headers += [f"{title} {n} - {s}" for s in subs]
    rows = []
    for c in cards:
        birthday = next((f.value for f in _by_kind(c, "date") if f.label.strip().lower() == "birthday"), "")
        labels = ["* myContacts", *[g for g in c.groups if g.strip()]] + (["* starred"] if c.favourite else [])
        row = [c.get("given"), c.get("additional"), c.get("family"), c.get("prefix"), c.get("suffix"),
               c.get("nickname"), c.get("company"), c.get("title"), c.get("department"), birthday,
               c.get("notes"), "", MULTI.join(labels)]
        for _title, kind, subs in spec:
            have = of(c, kind)
            for n in range(counts[kind]):
                f = have[n] if n < len(have) else None
                if f is None:
                    row += [""] * len(subs)
                elif kind == "address":
                    d = f.data
                    street = d.get("street", "")
                    first, _, rest = street.partition("\n")
                    row += [label_out(f.label), _address_text(f), first, d.get("city", ""), "",
                            d.get("region", ""), d.get("postcode", ""), d.get("country", ""), rest]
                else:
                    row += [label_out(f.label), f.value]
        rows.append(row)
    return headers, rows


_OUTLOOK_HEADERS = (
    ["Title", "First Name", "Middle Name", "Last Name", "Suffix", "Nickname", "E-mail Address", "E-mail 2 Address",
     "E-mail 3 Address", "Home Phone", "Home Phone 2", "Business Phone", "Business Phone 2", "Mobile Phone",
     "Car Phone", "Other Phone", "Primary Phone", "Company Main Phone", "Pager", "Business Fax", "Home Fax",
     "Other Fax", "IMAddress", "Job Title", "Department", "Company"]
    + [f"{block} {part}" for block in ("Business", "Home", "Other")
       for part in ("Street", "City", "State", "Postal Code", "Country/Region")]
    + ["Web Page", "Personal Web Page", "Spouse", "Manager's Name", "Assistant's Name", "Birthday", "Anniversary",
       "Notes", "Categories"])

_OUTLOOK_PARTS = {"Street": "street", "City": "city", "State": "region", "Postal Code": "postcode",
                  "Country/Region": "country"}


def _outlook_table(cards: list[Card]) -> tuple[list[str], list[list[str]]]:
    col = {h: i for i, h in enumerate(_OUTLOOK_HEADERS)}
    rows = []
    for c in cards:
        row = [""] * len(_OUTLOOK_HEADERS)

        def put(name: str, value: str) -> None:
            row[col[name]] = value

        for name, header in (("prefix", "Title"), ("given", "First Name"), ("additional", "Middle Name"),
                             ("family", "Last Name"), ("suffix", "Suffix"), ("nickname", "Nickname"),
                             ("title", "Job Title"), ("department", "Department"), ("company", "Company"),
                             ("notes", "Notes")):
            put(header, c.get(name))
        for slot, f in zip(("E-mail Address", "E-mail 2 Address", "E-mail 3 Address"), _by_kind(c, "email")):
            put(slot, f.value)
        for f in _by_kind(c, "phone"):
            wanted = f.label.strip().lower()
            for slot in L.OUTLOOK_PHONE_OUT.get(wanted, ["Other Phone"]):
                if not row[col[slot]]:
                    put(slot, f.value)
                    break
        used: set[str] = set()
        for f in _by_kind(c, "address"):
            lab = f.label.strip().lower()
            block = "Business" if lab == "work" else "Home" if lab == "home" else "Other"
            if block in used:
                block = "Other"
            if block in used:
                continue
            used.add(block)
            d = f.data or {"street": f.value}
            for part, key in _OUTLOOK_PARTS.items():
                put(f"{block} {part}", d.get(key, ""))
        for slot, f in zip(("Web Page", "Personal Web Page"), _by_kind(c, "url")):
            put(slot, f.value)
        ims = _by_kind(c, "im")
        if ims:
            put("IMAddress", ims[0].value)
        for f in _by_kind(c, "date"):
            lab = f.label.strip().lower()
            if lab in ("birthday", "anniversary") and not row[col[lab.capitalize()]]:
                put(lab.capitalize(), f.value)
        for f in _by_kind(c, "related"):
            slot = {"spouse": "Spouse", "manager": "Manager's Name", "assistant": "Assistant's Name"}.get(
                f.label.strip().lower())
            if slot and not row[col[slot]]:
                put(slot, f.value)
        put("Categories", ";".join(g for g in c.groups if g.strip()))
        rows.append(row)
    return list(_OUTLOOK_HEADERS), rows


def csv_table(cards: Iterable[Card], layout: str = "plain") -> tuple[list[str], list[list[str]]]:
    cards = list(cards)
    if layout == "plain":
        return _plain_table(cards)
    if layout == "google":
        return _google_table(cards)
    if layout == "outlook":
        return _outlook_table(cards)
    raise ValueError(f"The CSV layout {layout!r} is not known. Choose google, outlook or plain.")


def export_csv(cards: Iterable[Card], layout: str = "plain") -> str:
    """CSV text in one of the three layouts, comma-separated, CRLF line ends.
    Written to disk with `write_csv`, which adds the byte-order mark Excel needs."""
    headers, rows = csv_table(cards, layout)
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow(headers)
    w.writerows(rows)
    return out.getvalue()


def write_csv(cards: Iterable[Card], path, layout: str = "plain") -> Path:
    path = Path(path)
    text = export_csv(cards, layout)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:            # BOM: Excel reads UTF-8 only with it
        fh.write(text)
    return path


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

_XLSX_CELL_MAX = 32767


def export_xlsx(cards: Iterable[Card], path) -> Path:
    """One sheet, the plain layout without the photo. Text is stored as text: a note
    that starts with `=` is shown, never run."""
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # noqa: PLC0415
    from openpyxl.styles import Font  # noqa: PLC0415

    headers, rows = _plain_table(list(cards), with_photo=False)
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", v)[:_XLSX_CELL_MAX] for v in row])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    path = Path(path)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

_A4 = (210.0, 297.0)


def _top(page_h: float, y_mm: float, h_mm: float) -> float:
    """Points from the page bottom for a box whose top is `y_mm` from the page top."""
    return P.mm(page_h - y_mm - h_mm)


def export_pdf_list(cards: Iterable[Card], path, title: str = "Contacts") -> PdfLayout:
    """A printable list: name, company, email and phone, one row each, A4 portrait."""
    cards = list(cards)
    if not cards:
        raise NothingToPrint("There are no contacts to print. Choose some contacts, or clear the filter, and try again.")
    regular, bold, _ = P.fonts()
    W, H = _A4
    margin, row_h, head_h = 15.0, 6.5, 7.0
    cols = [("Name", 48.0), ("Company", 38.0), ("Email", 52.0), ("Phone", 42.0)]
    width = sum(w for _, w in cols)
    foot = H - 12.0
    pages: list[list[int]] = []
    cur: list[int] = []
    y_start = margin + 12.0
    y = y_start
    for i in range(len(cards)):
        if y + head_h + row_h > foot and cur:
            pages.append(cur)
            cur, y = [], margin
        if not cur:
            y = (y_start if not pages else margin) + head_h
        cur.append(i)
        y += row_h
    pages.append(cur)

    lay = PdfLayout(str(path), _A4, len(pages))
    c = P.new_canvas(path, _A4, title)
    for pno, idxs in enumerate(pages, 1):
        y = margin
        if pno == 1:
            c.setFont(bold, 14)
            c.drawString(P.mm(margin), _top(H, y, 6), P.safe(title))
            c.setFont(regular, 8)
            c.drawRightString(P.mm(margin + width), _top(H, y, 6), f"{len(cards)} contacts")
            y = y_start
        c.setFillGray(0.92)
        c.rect(P.mm(margin), _top(H, y, head_h), P.mm(width), P.mm(head_h), stroke=0, fill=1)
        c.setFillGray(0)
        c.setFont(bold, 9)
        x = margin
        for name, w in cols:
            c.drawString(P.mm(x + 1.5), _top(H, y, head_h) + P.mm(2.2), name)
            x += w
        y += head_h
        for i in idxs:
            card = cards[i]
            values = [display_name(card), card.get("company"), P.first_value(card, "email"),
                      P.first_value(card, "phone")]
            c.setFont(regular, 9)
            x = margin
            for (_n, w), v in zip(cols, values):
                c.drawString(P.mm(x + 1.5), _top(H, y, row_h) + P.mm(2.0), P.fit_text(v, regular, 9, P.mm(w - 3)))
                x += w
            c.setStrokeGray(0.85)
            c.setLineWidth(0.3)
            c.line(P.mm(margin), _top(H, y, row_h), P.mm(margin + width), _top(H, y, row_h))
            lay.placements.append(Placement(pno, i, margin, y, width, row_h, "row"))
            y += row_h
        c.setFont(regular, 8)
        c.setFillGray(0.4)
        c.drawCentredString(P.mm(W / 2), P.mm(8), f"Page {pno} of {len(pages)}")
        c.setFillGray(0)
        c.showPage()
    c.save()
    return lay


def export_pdf_cards(cards: Iterable[Card], path, page: str = "A4") -> PdfLayout:
    """Business-card grid, 85 x 55 mm cards, two across and five down: photo or
    initials, name, title, company, email, phone. Thin outlines are the cutting guide."""
    cards = list(cards)
    if not cards:
        raise NothingToPrint("There are no contacts to print. Choose some contacts, or clear the filter, and try again.")
    regular, bold, _ = P.fonts()
    size = P.page_size_mm(page)
    W, H = size
    cw, ch, ncol, nrow = 85.0, 55.0, 2, 5
    if W < cw * ncol or H < ch * nrow:
        ncol, nrow = int(W // cw), int(H // ch)          # Letter is 215.9 x 279.4: still 2 x 5
    left, top = (W - cw * ncol) / 2, (H - ch * nrow) / 2
    per_page = ncol * nrow
    lay = PdfLayout(str(path), size, (len(cards) + per_page - 1) // per_page)
    c = P.new_canvas(path, size, "Contact cards")
    for i, card in enumerate(cards):
        pno, slot = divmod(i, per_page)
        if slot == 0 and i:
            c.showPage()
        row, col = divmod(slot, ncol)
        x, y = left + col * cw, top + row * ch
        c.setStrokeGray(0.7)
        c.setLineWidth(0.25)
        c.rect(P.mm(x), _top(H, y, ch), P.mm(cw), P.mm(ch), stroke=1, fill=0)
        P.draw_photo_or_initials(c, card, P.mm(x + 5), _top(H, y + 5, 20), P.mm(20), lay.warnings)
        tx, tw = x + 29, cw - 29 - 4
        lines = [(display_name(card), bold, 11), (card.get("title"), regular, 8.5), (card.get("company"), regular, 8.5),
                 (P.first_value(card, "email"), regular, 8), (P.first_value(card, "phone"), regular, 8)]
        ty = y + 11.0
        for text, font, fsize in lines:
            if text:
                c.setFont(font, fsize)
                c.drawString(P.mm(tx), P.mm(H - ty), P.fit_text(text, font, fsize, P.mm(tw)))
                ty += fsize * 0.42 + 3.0
        lay.placements.append(Placement(pno + 1, i, x, y, cw, ch, "card", slot + 1))
    c.save()
    return lay
